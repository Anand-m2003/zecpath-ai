from docx import Document

document = Document()

document.add_heading("JOHN DOE", level=1)

document.add_heading("PROFESSIONAL SUMMARY", level=2)
document.add_paragraph(
    "Python Developer with experience in Flask, Django, REST APIs and SQL."
)

document.add_heading("TECHNICAL SKILLS", level=2)
document.add_paragraph("Python")
document.add_paragraph("Flask")
document.add_paragraph("Django")
document.add_paragraph("SQL")
document.add_paragraph("Git")
document.add_paragraph("REST API")

document.add_heading("EXPERIENCE", level=2)
document.add_paragraph("Python Developer - ABC Technologies")
document.add_paragraph("2023 - 2025")
document.add_paragraph("• Developed Flask-based web applications")
document.add_paragraph("• Created REST APIs")
document.add_paragraph("• Worked with SQL databases")

document.add_heading("EDUCATION", level=2)
document.add_paragraph("Bachelor of Computer Applications")
document.add_paragraph("ABC University")
document.add_paragraph("2023")

document.add_heading("PROJECTS", level=2)
document.add_paragraph("AI Resume Screening System")
document.add_paragraph(
    "Developed a system to extract and analyze resume information."
)
document.add_paragraph("Technologies: Python, Flask, NLP")

document.add_heading("CERTIFICATIONS", level=2)
document.add_paragraph("Python Programming Certification")

document.save("test_resume.docx")

print("DOCX test resume created successfully.")