from docx import Document


def extract_text_from_docx(file_path):
    """
    Extract text from paragraphs and tables in a DOCX resume.

    Args:
        file_path: Path to the DOCX resume.

    Returns:
        Extracted text as a string.
    """
    document = Document(file_path)

    content = []

    # Extract paragraphs
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            content.append(paragraph.text.strip())

    # Extract tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                content.append(" | ".join(row_text))

    return "\n".join(content)